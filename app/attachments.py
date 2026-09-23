import base64
import io
import zipfile
from pathlib import Path

MAX_BYTES=8*1024*1024


def extract(filename, body):
    if not body or len(body)>MAX_BYTES:
        raise ValueError("Файл должен быть непустым и не больше 8 МБ.")
    extension=Path(filename).suffix.lower()
    if extension in {".xlsx",".docx"}:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            if sum(i.file_size for i in archive.infolist())>30*1024*1024 or len(archive.infolist())>3000:
                raise ValueError("Слишком большой документ после распаковки.")
    if extension in {".jpg",".jpeg",".png"}:
        from PIL import Image
        image=Image.open(io.BytesIO(body))
        image.thumbnail((1600,1600))
        output=io.BytesIO(); image.convert("RGB").save(output,format="JPEG",quality=85)
        return {"name":filename,"image":"data:image/jpeg;base64,"+base64.b64encode(output.getvalue()).decode(),"text":""}
    if extension==".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(body))
        if reader.is_encrypted or len(reader.pages)>30:
            raise ValueError("Поддерживаются незашифрованные PDF до 30 страниц.")
        text="\n".join(p.extract_text() or "" for p in reader.pages)
        if not text.strip():
            raise ValueError("В PDF нет текстового слоя. Прикрепите нужную страницу как JPEG.")
    elif extension==".docx":
        from docx import Document
        doc=Document(io.BytesIO(body))
        text="\n".join([p.text for p in doc.paragraphs]+[" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    elif extension==".xlsx":
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
        text="\n".join(" | ".join(str(c or "") for c in row[:16]) for sheet in list(workbook)[:5]
                       for row in sheet.iter_rows(max_row=150,values_only=True))
        workbook.close()
    elif extension==".xls":
        import xlrd
        workbook=xlrd.open_workbook(file_contents=body)
        text="\n".join(" | ".join(str(c) for c in sheet.row_values(i)[:16]) for sheet in workbook.sheets()[:5]
                       for i in range(min(150,sheet.nrows)))
    elif extension in {".txt",".csv"}:
        text=body.decode("utf-8-sig")
    else:
        raise ValueError("Поддерживаются PDF, Word (.docx), Excel (.xlsx/.xls), JPEG, PNG, TXT и CSV.")
    return {"name":filename,"text":text[:20000],"image":None}
